"""
MAARA DOCX generator
--------------------
Converts the Markdown report from the Report Writer Agent into a Word file that
follows the same template as the MAARA Draft Report (and maara_pdf.py):

  * Cover page with title, MAARA, report type, subtitle and Field/Details table
  * Header / footer on every page (with automatic page numbers)
  * Real Word heading styles (Heading 1/2/3) - works with Word's Navigation Pane
  * Navy tables with alternating rows and "Table N: ..." captions
  * Bullets, nested bullets, numbered lists, bold / italic

Template details (names, header text, etc.) are shared with maara_pdf.py,
so you only edit TEMPLATE SETTINGS in ONE place.

Requires:  pip install python-docx
Keep this file in the same folder as maara_pdf.py.
"""

import re
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Shared settings + text helpers (single source of truth)
from maara_pdf import (
    ACADEMIC_YEAR,
    COVER_SUBTITLE,
    DEPARTMENT,
    FACULTY_MENTOR,
    FOOTER_LEFT,
    HEADER_LEFT,
    HEADER_RIGHT,
    PROJECT_TITLE,
    REPORT_LABEL,
    SUBMITTED_BY,
    clean_text,
    extract_topic,
    strip_markdown,
)

FONT = "Arial"
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
GREY = RGBColor(0x44, 0x44, 0x44)
NAVY_HEX = "1F3864"
ZEBRA_HEX = "F2F4F7"
BORDER_HEX = "BFC5CE"

PAGE_W_CM = 21.0
MARGIN_CM = 1.8
CONTENT_W_CM = PAGE_W_CM - 2 * MARGIN_CM      # 17.4 cm


# ============================================================
# Low-level XML helpers
# ============================================================


_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
              "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
              "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
              "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
              "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
              "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
                "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]


def _insert_ordered(parent, child, order):
    """Insert child into parent respecting the OOXML schema sequence."""
    name = child.tag.split("}")[1]
    idx = order.index(name)
    for old in parent.findall(child.tag):
        parent.remove(old)
    for existing in list(parent):
        ename = existing.tag.split("}")[1]
        if ename in order and order.index(ename) > idx:
            existing.addprevious(child)
            return
    parent.append(child)


def _apply_rfonts(rpr, name=FONT):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]          # theme fonts would override Arial
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)   # tcPr: shd comes after tcW, fine


def _cell_margins(table, top=70, bottom=70, left=110, right=110):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _insert_ordered(tbl_pr, mar, _TBLPR_ORDER)


def _table_borders(table, hex_color=BORDER_HEX, size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), hex_color)
        borders.append(el)
    _insert_ordered(tbl_pr, borders, _TBLPR_ORDER)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _para_border(paragraph, side="bottom", hex_color=BORDER_HEX, size=6, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), hex_color)
    borders.append(el)
    _insert_ordered(p_pr, borders, _PPR_ORDER)


def _add_field(paragraph, instr, size=7.5, color=GREY):
    """Insert a Word field (e.g. PAGE) into a paragraph."""
    def _run():
        r = paragraph.add_run()
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _apply_rfonts(r._r.get_or_add_rPr())
        return r

    r1 = _run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r1._r.append(b)

    r2 = _run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._r.append(it)

    r3 = _run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)

    r4 = _run()
    r4.text = "1"

    r5 = _run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)


# ============================================================
# Inline markdown -> runs
# ============================================================

_INLINE_RE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<![\*\w])\*(?!\s).+?(?<!\s)\*(?![\*\w])|`.+?`)")


def add_inline(paragraph, text, size=None, color=None, bold=False, italic=False):
    text = clean_text(text)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"\[([^\]]+)\]\((?:[^)]+)\)", r"\1", text)

    for part in _INLINE_RE.split(text):
        if not part:
            continue
        b, i, code = bold, italic, False
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            part, b, i = part[3:-3], True, True
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, code = part[1:-1], True

        lines = part.split("\n")
        for n, seg in enumerate(lines):
            run = paragraph.add_run(seg)
            run.bold = True if b else None        # None = inherit from style (keeps headings bold)
            run.italic = True if i else None
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
            if code:
                run.font.name = "Courier New"
            else:
                _apply_rfonts(run._r.get_or_add_rPr())
            if n < len(lines) - 1:
                run.add_break(WD_BREAK.LINE)


# ============================================================
# Styles / page setup
# ============================================================

def _setup_styles(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN_CM)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.header_distance, sec.footer_distance = Cm(1.1), Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    _apply_rfonts(normal.element.get_or_add_rPr())
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    specs = {
        "Heading 1": (15, NAVY, 16, 6),
        "Heading 2": (12.5, NAVY, 12, 5),
        "Heading 3": (11, BLUE, 9, 4),
    }
    for name, (size, color, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = color
        _apply_rfonts(st.element.get_or_add_rPr())
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.keep_with_next = True
        pf.line_spacing = 1.1


def _clear_style_tabs(style):
    p_pr = style.element.get_or_add_pPr()
    for old in p_pr.findall(qn("w:tabs")):
        p_pr.remove(old)


def _setup_header_footer(doc):
    sec = doc.sections[0]
    right_tab = Cm(CONTENT_W_CM)
    _clear_style_tabs(doc.styles["Header"])
    _clear_style_tabs(doc.styles["Footer"])

    # Header
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_after = Pt(0)
    add_inline(hp, HEADER_LEFT.replace("&amp;", "&"), size=8, color=GREY)
    r = hp.add_run("\t" + HEADER_RIGHT.replace("&amp;", "&"))
    r.font.size, r.font.color.rgb = Pt(8), GREY
    _apply_rfonts(r._r.get_or_add_rPr())
    _para_border(hp, "bottom")

    # Footer
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.space_before = Pt(0)
    _para_border(fp, "top")
    add_inline(fp, FOOTER_LEFT, size=7.5, color=GREY)
    r = fp.add_run("\tPage ")
    r.font.size, r.font.color.rgb = Pt(7.5), GREY
    _apply_rfonts(r._r.get_or_add_rPr())
    _add_field(fp, "PAGE")


# ============================================================
# Tables
# ============================================================

def _column_widths(rows, total_cm=CONTENT_W_CM):
    cols = max(len(r) for r in rows)
    weights, mins = [], []
    for c in range(cols):
        texts = [strip_markdown(r[c]) if c < len(r) else "" for r in rows]
        longest = max(len(t) for t in texts)
        longest_word = max((len(w) for t in texts for w in t.split()), default=4)
        weights.append(min(max(longest, 12), 70) ** 0.8)
        mins.append((longest_word * 5.6 + 16) / 28.35)          # points -> cm
    total = sum(weights)
    widths = [total_cm * w / total for w in weights]
    for _ in range(5):
        deficit = sum(max(m - w, 0) for w, m in zip(widths, mins))
        if deficit <= 0:
            break
        big = [i for i, (w, m) in enumerate(zip(widths, mins)) if w > m]
        spare = sum(widths[i] - mins[i] for i in big) or 1
        widths = [m if w < m else w - deficit * (w - m) / spare
                  for w, m in zip(widths, mins)]
    scale = total_cm / sum(widths)
    return [w * scale for w in widths]


def add_table(doc, rows, widths_cm=None):
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    widths_cm = widths_cm or _column_widths(rows)

    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for c_idx, w in enumerate(widths_cm):
        table.columns[c_idx].width = Cm(w)
    _table_borders(table)
    _cell_margins(table)

    for r_idx, row in enumerate(rows):
        _cant_split(table.rows[r_idx])
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = Cm(widths_cm[c_idx])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if r_idx == 0:
                add_inline(p, text, size=8.8, color=RGBColor(255, 255, 255), bold=True)
                _shade(cell, NAVY_HEX)
            else:
                add_inline(p, text, size=8.6)
                if r_idx % 2 == 0:
                    _shade(cell, ZEBRA_HEX)
    _repeat_header(table.rows[0])
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    add_inline(p, text, size=8.5, color=GREY, italic=True)


# ============================================================
# Cover page
# ============================================================

def add_cover(doc, topic, report_date):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MULTI-AGENT AI RESEARCH ASSISTANT")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(22), NAVY
    _apply_rfonts(r._r.get_or_add_rPr())

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("MAARA")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15), BLUE
    _apply_rfonts(r._r.get_or_add_rPr())

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(REPORT_LABEL)
    r.bold, r.font.size = True, Pt(15)
    _apply_rfonts(r._r.get_or_add_rPr())

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run(COVER_SUBTITLE)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9.5), GREY
    _apply_rfonts(r._r.get_or_add_rPr())

    info = [
        ["Field", "Details"],
        ["Project Title", PROJECT_TITLE],
        ["Research Topic", topic],
        ["Report Type", REPORT_LABEL],
        ["Submitted by", SUBMITTED_BY],
        ["Department", DEPARTMENT],
        ["Faculty Mentor", FACULTY_MENTOR],
        ["Academic Year", ACADEMIC_YEAR],
        ["Date", report_date],
    ]
    # plain text in the info table (no markdown parsing surprises)
    info = [[c.replace("*", "") for c in row] for row in info]
    add_table(doc, info, widths_cm=[4.4, CONTENT_W_CM - 4.4])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("This report was generated from the research papers retrieved and analysed "
                  "by the Multi-Agent AI Research Assistant (MAARA).")
    r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), GREY
    _apply_rfonts(r._r.get_or_add_rPr())

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


# ============================================================
# Main
# ============================================================

_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET_RE = re.compile(r"^(\s*)([-*+\u2022\u25e6\u25cb\u25aa])\s+(.*)$")
_NUMBER_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_HR_RE = re.compile(r"^\s*([-*_])\1{2,}\s*$")


def _is_table_line(line):
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def _is_separator(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return any(cells) and all(re.fullmatch(r":?-{1,}:?", c) for c in cells if c != "")


def _split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _list_paragraph(doc, marker, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    left = Cm(0.9 + 0.8 * level)
    pf.left_indent = left
    pf.first_line_indent = Cm(-0.6)
    pf.space_after = Pt(3)
    pf.tab_stops.add_tab_stop(left)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(marker + "\t")
    _apply_rfonts(r._r.get_or_add_rPr())
    add_inline(p, text, size=9.6 if level else None)
    return p


def _flush(buf, doc):
    if buf:
        text = " ".join(buf).strip()
        if text:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, text)
        buf.clear()


def create_docx(report_text, filename="MAARA_Final_Report.docx", report_date=None):
    if not report_text or not str(report_text).strip():
        raise ValueError("Report text is empty. DOCX cannot be generated.")

    text = clean_text(str(report_text)).replace("\r\n", "\n")
    report_date = report_date or datetime.now().strftime("%B %Y")

    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    _setup_styles(doc)
    _setup_header_footer(doc)
    doc.core_properties.title = "MAARA Final Research Report"
    doc.core_properties.author = SUBMITTED_BY

    add_cover(doc, extract_topic(text), report_date)

    lines = text.split("\n")
    buf, last_heading, table_no, i = [], "", 0, 0

    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()

        if not s:
            _flush(buf, doc)
            i += 1
            continue

        if _HR_RE.match(s):
            _flush(buf, doc)
            i += 1
            continue

        if _is_table_line(s):
            _flush(buf, doc)
            rows = []
            while i < len(lines) and _is_table_line(lines[i].strip()):
                if not _is_separator(lines[i].strip()):
                    rows.append(_split_row(lines[i].strip()))
                i += 1
            if rows:
                table_no += 1
                add_table(doc, rows)
                cap = re.sub(r"^\d+(\.\d+)*\.?\s*", "", last_heading) or "Summary"
                add_caption(doc, f"Table {table_no}: {cap}")
            continue

        m = _HEADING_RE.match(line)
        if m:
            _flush(buf, doc)
            level = len(m.group(1))
            title = strip_markdown(m.group(2))
            if title and not (level == 1 and title.lower() in
                              {"maara", "research report", "final research report"}):
                last_heading = title
                style = "Heading 1" if level <= 2 else ("Heading 2" if level == 3 else "Heading 3")
                h = doc.add_paragraph(style=style)
                add_inline(h, title)
            i += 1
            continue

        m = _BULLET_RE.match(line)
        if m and not re.fullmatch(r"\*\*[^*]+\*\*:?", s):
            _flush(buf, doc)
            level = 1 if len(m.group(1).replace("\t", "    ")) >= 2 else 0
            _list_paragraph(doc, "\u2022" if level == 0 else "\u2013", m.group(3), level)
            i += 1
            continue

        m = _NUMBER_RE.match(line)
        if m:
            _flush(buf, doc)
            level = 1 if len(m.group(1).replace("\t", "    ")) >= 2 else 0
            _list_paragraph(doc, f"{m.group(2)}.", m.group(3), level)
            i += 1
            continue

        if re.fullmatch(r"\*\*[^*]+\*\*:?", s):
            _flush(buf, doc)
            h = doc.add_paragraph(style="Heading 3")
            add_inline(h, s)
            i += 1
            continue

        buf.append(s)
        i += 1

    _flush(buf, doc)
    doc.save(filename)
    return filename
