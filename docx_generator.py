from docx import Document
from docx.shared import Pt


def create_docx(report_text, filename="MAARA_Final_Report.docx"):

    doc = Document()

    title = doc.add_heading("MAARA", level=1)
    title.runs[0].font.size = Pt(22)

    doc.add_heading("Multi-Agent AI Research Assistant", level=2)

    doc.add_paragraph()

    for line in report_text.split("\n"):

        line = line.strip()

        if not line:
            continue

        # Main Headings
        if line.startswith("#"):

            clean = line.replace("#", "").strip()

            doc.add_heading(clean, level=2)

        else:

            doc.add_paragraph(line)

    doc.save(filename)

    return filename